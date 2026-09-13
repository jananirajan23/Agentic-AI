"""Unit tests for agents/report_aggregator.py."""

import os

from docx import Document

import agents.report_aggregator as agg


def test_compute_composite_score_uses_configured_weights():
    state = {"literature_score": 100, "architecture_score": 0, "alignment_score": 50}
    # weights: literature 0.30, architecture 0.35, alignment 0.35
    expected = round(100 * 0.30 + 0 * 0.35 + 50 * 0.35, 1)
    assert agg.compute_composite_score(state) == expected


def test_generate_suggestions_without_groq_key_falls_back_to_flags(monkeypatch):
    monkeypatch.setattr(agg.config, "GROQ_API_KEY", "")
    flags = ["Issue A", "Issue B"]
    assert agg.generate_suggestions({}, flags) == flags


def test_generate_suggestions_no_flags(monkeypatch):
    monkeypatch.setattr(agg.config, "GROQ_API_KEY", "")
    result = agg.generate_suggestions({}, [])
    assert result == ["No issues were flagged."]


def test_build_report_writes_readable_docx(tmp_path):
    state = {
        "composite_score": 72.5,
        "literature_score": 60,
        "architecture_score": 80,
        "alignment_score": 75,
        "citation_resolution_rate": 0.5,
        "embedding_alignment_score": 65.0,
        "parse_warnings": ["A warning."],
        "all_flags": ["Flag one.", "Flag two."],
        "research_gap_summary": "A summary of the gap.",
        "suggestions": ["Do better."],
    }
    output_path = os.path.join(tmp_path, "draft.docx")
    agg.build_report(state, output_path)

    assert os.path.exists(output_path)
    doc = Document(output_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "AI Draft" in full_text
    assert "72.5" in full_text
    assert "Flag one." in full_text
    assert "Do better." in full_text
    assert "has not yet been reviewed by a faculty member" in full_text


def test_build_final_report_marks_approval_and_filters_content(tmp_path):
    state = {
        "composite_score": 72.5,
        "literature_score": 60,
        "architecture_score": 80,
        "alignment_score": 75,
        "citation_resolution_rate": 0.5,
        "embedding_alignment_score": 65.0,
        "parse_warnings": [],
    }
    output_path = os.path.join(tmp_path, "final.docx")
    agg.build_final_report(
        state,
        output_path,
        included_flags=["Kept flag."],
        included_suggestions=["Kept suggestion."],
        research_gap_summary="Edited gap summary.",
        approved_by="Dr. Reviewer",
        approved_at="2026-01-01 10:00 UTC",
    )

    doc = Document(output_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Faculty-Approved" in full_text
    assert "Dr. Reviewer" in full_text
    assert "Kept flag." in full_text
    assert "Kept suggestion." in full_text
    assert "Edited gap summary." in full_text


def test_run_produces_final_state_keys(monkeypatch, tmp_path):
    monkeypatch.setattr(agg.config, "GROQ_API_KEY", "")
    state = {
        "literature_score": 50,
        "architecture_score": 50,
        "alignment_score": 50,
        "literature_flags": ["L flag"],
        "architecture_flags": [],
        "alignment_flags": [],
    }
    output_path = os.path.join(tmp_path, "review.docx")
    result = agg.run(state, output_path=output_path)

    assert result["composite_score"] == 50.0
    assert result["all_flags"] == ["L flag"]
    assert os.path.exists(result["report_path"])
