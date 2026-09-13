"""
Report Aggregator Agent.

Combines the three agent scores into a composite score, turns the raw
flags into constructive student-facing suggestions, and writes the whole
review out as a downloadable .docx.
"""

import json
import re
from datetime import datetime
from typing import Dict, List

from docx import Document
from docx.shared import Pt, RGBColor
from groq import Groq

import config


def _strip_code_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def compute_composite_score(state: dict) -> float:
    weights = config.SCORE_WEIGHTS
    total = (
        state.get("literature_score", 0) * weights["literature_score"]
        + state.get("architecture_score", 0) * weights["architecture_score"]
        + state.get("alignment_score", 0) * weights["alignment_score"]
    )
    return round(total, 1)


def generate_suggestions(state: dict, all_flags: List[str]) -> List[str]:
    if not config.GROQ_API_KEY:
        return list(all_flags) if all_flags else ["No issues were flagged."]

    if not all_flags:
        return ["No issues were flagged — the report is in good shape across all three checks."]

    client = Groq(api_key=config.GROQ_API_KEY)
    flags_text = "\n".join(f"- {f}" for f in all_flags)

    prompt = f"""You are helping a student improve their capstone project report before faculty evaluation. Below are issues an automated pre-review flagged.

Flagged issues:
{flags_text}

Turn these into 3-6 constructive, specific, student-facing improvement suggestions. Be direct but supportive in tone — this is meant to help the student, not discourage them. Group related flags into a single suggestion where it makes sense; don't just restate each flag verbatim.

Respond with STRICT JSON only, no markdown fences: a JSON array of strings, e.g. ["<suggestion 1>", "<suggestion 2>", ...]."""

    try:
        resp = client.chat.completions.create(
            model=config.GROQ_TEXT_MODEL,
            messages=[{"role": "user", "content": prompt}],
            max_tokens=config.MAX_TOKENS,
            temperature=0.4,
        )
        content = _strip_code_fences(resp.choices[0].message.content or "")
        data = json.loads(content)
        if isinstance(data, list) and data:
            return [str(s) for s in data]
        return list(all_flags)
    except Exception:
        return list(all_flags)


def _add_bold_heading(doc: Document, text: str, size: int = 14):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    return p


def _add_bulleted_list(doc: Document, items: List[str]):
    if not items:
        doc.add_paragraph("No issues flagged.")
        return
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def build_report(state: dict, output_path: str, approval: Dict | None = None) -> str:
    """Builds the review .docx. `approval`, when given, marks this as the
    faculty-approved (student-facing) version rather than the AI draft —
    {"approved_by": str, "approved_at": str}. Ch.3.3/Ch.4.1 of the
    requirements treat this distinction as the core human-in-the-loop
    requirement: AI findings must not reach students without it."""
    doc = Document()

    if approval:
        _add_bold_heading(doc, "Reviewer-Lens: Faculty-Approved Pre-Review Report", size=20)
        meta = doc.add_paragraph()
        meta.add_run(
            f"Approved by {approval.get('approved_by', 'faculty reviewer')} on {approval.get('approved_at', '')}"
        ).italic = True
        disclaimer_text = (
            "This report has been reviewed and approved by the faculty member named above. Items the "
            "reviewer excluded during approval are not included below."
        )
    else:
        _add_bold_heading(doc, "Reviewer-Lens: Capstone Pre-Review Report (AI Draft)", size=20)
        ts = datetime.now().strftime("%Y-%m-%d %H:%M")
        meta = doc.add_paragraph()
        meta.add_run(f"Generated: {ts}").italic = True
        disclaimer_text = (
            "This automated report is a pre-review aid intended to support faculty evaluation. "
            "It does not replace faculty judgment, has not yet been reviewed by a faculty member, "
            "and should not be shared with students as final feedback."
        )

    disclaimer = doc.add_paragraph()
    disclaimer_run = disclaimer.add_run(disclaimer_text)
    disclaimer_run.italic = True
    disclaimer_run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.add_paragraph()

    composite = state.get("composite_score", 0)
    _add_bold_heading(doc, f"Composite Score: {composite} / 100", size=16)

    doc.add_paragraph()
    _add_bold_heading(doc, "Score Breakdown", size=13)
    weights = config.SCORE_WEIGHTS
    doc.add_paragraph(
        f"Literature Survey: {state.get('literature_score', 0)} / 100  "
        f"(weight {int(weights['literature_score'] * 100)}%)",
        style="List Bullet",
    )
    doc.add_paragraph(
        f"Architecture Structure: {state.get('architecture_score', 0)} / 100  "
        f"(weight {int(weights['architecture_score'] * 100)}%)",
        style="List Bullet",
    )
    doc.add_paragraph(
        f"Alignment: {state.get('alignment_score', 0)} / 100  "
        f"(weight {int(weights['alignment_score'] * 100)}%)",
        style="List Bullet",
    )

    doc.add_paragraph()
    _add_bold_heading(doc, "Measured Retrieval & Alignment Metrics", size=13)
    doc.add_paragraph(
        f"Citation resolution rate (Semantic Scholar): "
        f"{state.get('citation_resolution_rate', 0) * 100:.0f}% of citations found in the report were "
        "resolved against Semantic Scholar.",
        style="List Bullet",
    )
    doc.add_paragraph(
        f"Embedding-based semantic alignment score: {state.get('embedding_alignment_score', 0)} / 100 "
        "(deterministic cosine-similarity match between diagram components and module-description "
        "phrases, independent of LLM judgment).",
        style="List Bullet",
    )

    warnings = state.get("parse_warnings") or []
    if warnings:
        doc.add_paragraph()
        _add_bold_heading(doc, "Parsing Warnings", size=13)
        _add_bulleted_list(doc, warnings)

    doc.add_paragraph()
    _add_bold_heading(doc, "Flagged Issues", size=13)
    _add_bulleted_list(doc, state.get("all_flags") or [])

    gap_summary = state.get("research_gap_summary") or ""
    if gap_summary.strip():
        doc.add_paragraph()
        _add_bold_heading(doc, "Research Gap Summary", size=13)
        doc.add_paragraph(gap_summary)

    doc.add_paragraph()
    _add_bold_heading(doc, "Improvement Suggestions", size=13)
    _add_bulleted_list(doc, state.get("suggestions") or [])

    doc.save(output_path)
    return output_path


def build_final_report(
    state: dict,
    output_path: str,
    included_flags: List[str],
    included_suggestions: List[str],
    research_gap_summary: str,
    approved_by: str,
    approved_at: str,
) -> str:
    """Builds the faculty-approved report: the same layout as the AI
    draft, but restricted to the flags/suggestions the reviewer kept, with
    a reviewer-editable gap summary and an approval byline."""
    final_state = dict(state)
    final_state["all_flags"] = included_flags
    final_state["suggestions"] = included_suggestions
    final_state["research_gap_summary"] = research_gap_summary
    return build_report(final_state, output_path, approval={"approved_by": approved_by, "approved_at": approved_at})


def run(state: dict, output_path: str = None) -> dict:
    """LangGraph node entry point. Returns only the keys this agent adds."""
    all_flags: List[str] = []
    all_flags.extend(state.get("literature_flags") or [])
    all_flags.extend(state.get("architecture_flags") or [])
    all_flags.extend(state.get("alignment_flags") or [])

    composite_score = compute_composite_score(state)
    suggestions = generate_suggestions(state, all_flags)

    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d%H%M%S")
        import os

        output_path = os.path.join(config.REPORTS_DIR, f"review_{ts}.docx")

    report_state = dict(state)
    report_state["composite_score"] = composite_score
    report_state["all_flags"] = all_flags
    report_state["suggestions"] = suggestions

    report_path = build_report(report_state, output_path)

    return {
        "composite_score": composite_score,
        "all_flags": all_flags,
        "suggestions": suggestions,
        "report_path": report_path,
    }
